# -*- coding: utf-8 -*-
"""
jetgen - build Jet Reports workbooks (.xlsx with NL/NF/NP formulas) from a YAML spec.

The file is written the way Excel 365 stores Jet formulas: `_xll.NL(...)` as a dynamic-array
formula (<c cm="1"><f t="array">) plus xl/metadata.xml. Verified with Jet Reports 23.4.0.3181,
Excel 365 (de-CH), Microsoft Dynamics NAV 2017 (Oct 2026).

Usage:
  python jetgen.py <spec.yaml> <out.xlsx> [--profile profile.yaml]   check, then build
  python jetgen.py --check <spec.yaml> [--profile profile.yaml]      check only

Not affiliated with insightsoftware. "Jet Reports" is a trademark of its owner.
"""
import argparse
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.packaging.custom import BoolProperty, StringProperty
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula

JET_FUNCTIONS = ("NL", "NF", "NP", "GL", "NFL")

# Profile = environment/language settings. Override per company with --profile or build(profile=...).
DEFAULT_PROFILE = {
    "keywords": {                       # Jet keywords as your Jet UI language writes them
        "auto": "Auto+Hide", "values": "+Values", "hide_sheet": "+HideSheet",
        "hide": "Hide", "show": "Show", "option": "Option",
    },
    "options_sheet": "Options",
    "option_tags": ["Title", "Value", "Lookup", "Tooltip"],
    "version_sheet": "Version",
    "version_headers": ["Date", "Requested by", "Changed by", "Description", "Version"],
    "company_key": "Company=",          # e.g. "Mandant=" in a German Jet
    "company_keys_accepted": ["Company="],
    "global_tables": ["2000000006"],    # tables queried without a company (Company table)
    "require_table_numbers": False,     # True: tables must be written as "27 Item" (number + name)
    "header_color": "1F4E79",
    "localized_function_names": [],     # e.g. German WENN, SUMME -> flagged (file format needs English)
    "function_literals": None,          # optional docProps value "Jet Reports Function Literals"
    "number_formats": {"int": "#,##0", "dec": "#,##0.00", "date": "DD.MM.YYYY",
                       "pct": "0.0%", "text": "@", "datetime": "DD.MM.YYYY hh:mm"},
}
EXCEL_WORDS = {"TRUE", "FALSE"}

DYN_METADATA = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
    'xmlns:xda="http://schemas.microsoft.com/office/spreadsheetml/2017/dynamicarray">'
    '<metadataTypes count="1"><metadataType name="XLDAPR" minSupportedVersion="120000" copy="1" '
    'pasteAll="1" pasteValues="1" merge="1" splitFirst="1" rowColShift="1" clearFormats="1" '
    'clearComments="1" assign="1" coerce="1" cellMeta="1"/></metadataTypes>'
    '<futureMetadata name="XLDAPR" count="1"><bk><extLst><ext uri="{bdbb8cdc-fa1e-496e-a857-3c3f30c029c3}">'
    '<xda:dynamicArrayProperties fDynamic="1" fCollapsed="0"/></ext></extLst></bk></futureMetadata>'
    '<cellMetadata count="1"><bk><rc t="1" v="0"/></bk></cellMetadata></metadata>'
)

RE_PLACEHOLDER = re.compile(r"\[\[([^\]]+)\]\]")
RE_JET = re.compile(r"(?<![\w.])(%s)\(" % "|".join(JET_FUNCTIONS))


def merge_profile(override=None):
    p = {k: (dict(v) if isinstance(v, dict) else v) for k, v in DEFAULT_PROFILE.items()}
    for k, v in (override or {}).items():
        if isinstance(v, dict) and isinstance(p.get(k), dict):
            p[k].update(v)
        else:
            p[k] = v
    if p["company_key"] not in p["company_keys_accepted"]:
        p["company_keys_accepted"] = list(p["company_keys_accepted"]) + [p["company_key"]]
    return p


# ---------------------------------------------------------------- helpers
def _strip_strings(formula):
    return re.sub(r'"(?:[^"]|"")*"', '""', formula)


def _is_jet(formula):
    return bool(RE_JET.search(_strip_strings(formula)))


def _with_prefix(formula):
    """NL( -> _xll.NL( outside string literals."""
    parts = re.split(r'("(?:[^"]|"")*")', formula)
    return "".join(t if t.startswith('"') else RE_JET.sub(r"_xll.\1(", t) for t in parts)


def _column_plan(table):
    start = column_index_from_string(table.get("start_column", "C"))
    return {c["header"]: get_column_letter(start + i) for i, c in enumerate(table["columns"])}


def _fill_placeholders(formula, plan, row):
    return RE_PLACEHOLDER.sub(lambda m: f"${plan[m.group(1)]}{row}", formula)


def _nl_args(formula):
    """Rough argument lists of every NL(...) call (for checks)."""
    result = []
    for m in re.finditer(r"(?<![\w.])NL\(", formula):
        i, depth, args, cur, in_str = m.end(), 1, [], "", False
        while i < len(formula) and depth:
            c = formula[i]
            if c == '"':
                in_str = not in_str
            elif not in_str and c in "({":
                depth += 1
            elif not in_str and c in ")}":
                depth -= 1
                if depth == 0:
                    break
            if not in_str and c == "," and depth == 1:
                args.append(cur.strip())
                cur = ""
            else:
                cur += c
            i += 1
        args.append(cur.strip())
        result.append(args)
    return result


# ---------------------------------------------------------------- checks
def _check_formula(formula, where, names, plan, prof):
    errors = []
    raw = _strip_strings(formula)
    if raw.count("(") != raw.count(")") or formula.count('"') % 2:
        errors.append(f"{where}: unbalanced parentheses or quotes")
    for word in prof["localized_function_names"]:
        if re.search(r"(?<![\w.])%s\(" % re.escape(word), raw):
            errors.append(f"{where}: function {word} - the file format needs English function names "
                          f"and comma separators (IF, SUM, ...)")
    for ph in RE_PLACEHOLDER.findall(formula):
        if plan is None or ph not in plan:
            errors.append(f"{where}: placeholder [[{ph}]] does not match a table column header")
    no_ph = RE_PLACEHOLDER.sub("X1", raw)
    for m in re.finditer(r"(?<![\w.$!'])([A-Za-zÄÖÜäöü_][\w.]*)(?![\w.(!])", no_ph):
        name = m.group(1)
        if re.fullmatch(r"[A-Za-z]{1,3}\d+", name) or name.upper() in EXCEL_WORDS or name in names:
            continue
        errors.append(f"{where}: name '{name}' is not defined (add it as option or under 'names')")
    for args in _nl_args(formula):
        if len(args) < 2:
            continue
        table = args[1]
        if table.startswith('"') and not re.fullmatch(r'"\d+ [^"]+"', table) and prof.get("require_table_numbers"):
            errors.append(f"{where}: NL table {table} without table number (pattern \"27 Item\")")
        tno = re.match(r'"(\d+)', table)
        has_company = any(a.strip('"') in prof["company_keys_accepted"] for a in args)
        rows_from_np = args[0].strip('"').lower().startswith(("rows", "zeilen")) and len(args) == 2
        if not has_company and not rows_from_np and not (tno and tno.group(1) in prof["global_tables"]):
            errors.append(f"{where}: NL without \"{prof['company_key']}\" - every query needs a company")
    return errors


def check(spec, profile=None):
    """Return a list of error strings (empty = OK)."""
    prof = merge_profile(profile)
    errors = []
    names = {o["name"] for o in spec.get("options", [])} | set(spec.get("names", {}))
    for o in spec.get("options", []):
        if o.get("lookup"):
            errors += _check_formula(o["lookup"], f"option {o['name']}/lookup", names, None, prof)
    for s in spec.get("sheets", []):
        plan = _column_plan(s["table"]) if s.get("table") else None
        for c in s.get("cells", []):
            if c.get("formula"):
                errors += _check_formula(c["formula"], f"{s['name']}!{c['cell']}", names, plan, prof)
        if s.get("table"):
            t = s["table"]
            for col in t["columns"]:
                errors += _check_formula(col["formula"], f"{s['name']}/{col['header']}", names, plan, prof)
            if t.get("hide_row_if"):
                errors += _check_formula(t["hide_row_if"], f"{s['name']}/hide_row_if", names, plan, prof)
            if not _is_jet(t["columns"][0]["formula"]):
                errors.append(f"{s['name']}: first table column must be the NL(\"Rows\",...) replicator")
    return errors


# ---------------------------------------------------------------- build
def _put(ws, ref, content, fmt, prof):
    if isinstance(content, str) and content.startswith("="):
        f = content[1:]
        ws[ref] = ArrayFormula(ref, "=" + _with_prefix(f)) if _is_jet(f) else content
    else:
        ws[ref] = content
    if fmt:
        ws[ref].number_format = prof["number_formats"].get(fmt, fmt)


def _header_style(cell, prof):
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=prof["header_color"])
    cell.alignment = Alignment(vertical="center", wrap_text=True)


def build(spec, out, profile=None):
    errors = check(spec, profile)
    if errors:
        raise ValueError("spec check failed:\n- " + "\n- ".join(errors))
    prof = merge_profile(profile)
    kw = prof["keywords"]
    values = kw["values"] if spec.get("mode") == "values" else ""
    wb = Workbook()

    # options sheet (Jet Report Options)
    opt = wb.active
    opt.title = prof["options_sheet"]
    opt["A1"] = kw["auto"] + kw["hide_sheet"]
    for col, tag in zip("CDEF", prof["option_tags"]):
        opt[f"{col}1"] = tag
        _header_style(opt[f"{col}1"], prof)
    for i, o in enumerate(spec.get("options", []), start=3):
        opt[f"A{i}"], opt[f"C{i}"] = kw["option"], o.get("title", o["name"])
        _put(opt, f"D{i}", o.get("value"), o.get("format"), prof)
        if o.get("lookup"):
            _put(opt, f"E{i}", "=" + o["lookup"], None, prof)
        if o.get("tooltip"):
            opt[f"F{i}"] = o["tooltip"]
        sheet_ref = opt.title if re.fullmatch(r"\w+", opt.title) else f"'{opt.title}'"
        wb.defined_names[o["name"]] = DefinedName(o["name"], attr_text=f"{sheet_ref}!$D${i}")
    for n, target in spec.get("names", {}).items():
        wb.defined_names[n] = DefinedName(n, attr_text=target)
    for col, w in zip("ABCDEF", [4, 2, 28, 24, 30, 50]):
        opt.column_dimensions[col].width = w

    # report sheets
    for s in spec["sheets"]:
        ws = wb.create_sheet(s["name"])
        ws["A1"] = kw["auto"] + values + (kw["hide_sheet"] if s.get("hide_sheet") else "")
        if s.get("title"):
            _put(ws, "B3", s["title"], None, prof)
            ws["B3"].font = Font(bold=True, size=14, color=prof["header_color"])
        t = s.get("table")
        plan = _column_plan(t) if t else {}
        for c in s.get("cells", []):
            content = c.get("value") if "value" in c else "=" + c["formula"]
            _put(ws, c["cell"], content, c.get("format"), prof)
        if t:
            hr = t.get("header_row", 6)
            dr = hr + 1
            for col in t["columns"]:
                letter = plan[col["header"]]
                ws[f"{letter}{hr}"] = col["header"]
                _header_style(ws[f"{letter}{hr}"], prof)
                _put(ws, f"{letter}{dr}", "=" + _fill_placeholders(col["formula"], plan, dr), col.get("format"), prof)
                ws.column_dimensions[letter].width = col.get("width", 16)
                if col.get("hide"):
                    ws[f"{letter}1"] = kw["hide"]
            if t.get("hide_row_if"):
                cond = _fill_placeholders(t["hide_row_if"], plan, dr)
                _put(ws, f"A{dr}", f'=IF({cond},"{kw["hide"]}","{kw["show"]}")', None, prof)
            ws.freeze_panes = f"A{dr}"

    # version sheet
    v = wb.create_sheet(prof["version_sheet"])
    v["A1"] = kw["auto"] + kw["hide_sheet"]
    for col, txt in zip("CDEFG", prof["version_headers"]):
        v[f"{col}3"] = txt
        _header_style(v[f"{col}3"], prof)
    for i, e in enumerate(spec.get("version", []), start=4):
        for col, key in zip("CDEFG", ["date", "requested_by", "changed_by", "description", "version"]):
            v[f"{col}{i}"] = e.get(key)
        v[f"C{i}"].number_format = prof["number_formats"]["date"]
    for col, w in zip("CDEFG", [12, 18, 18, 60, 9]):
        v.column_dimensions[col].width = w

    if prof.get("function_literals"):
        wb.custom_doc_props.append(StringProperty(name="Jet Reports Function Literals",
                                                  value=prof["function_literals"]))
    wb.custom_doc_props.append(BoolProperty(name="Jet Reports Design Mode Active", value=True))
    out = Path(out)
    wb.save(out)
    _make_dynamic(out)
    return out


def _make_dynamic(path):
    """openpyxl cannot write cm="1": mark array formulas as dynamic arrays (Excel 365 format)."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                x = re.sub(r'<c r="([A-Z]+\d+)"([^>]*)>(<f t="array")',
                           lambda m: f'<c r="{m.group(1)}"{m.group(2)} cm="1">{m.group(3)}', data.decode("utf-8"))
                data = x.encode("utf-8")
            elif item.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", b'<Override PartName="/xl/metadata.xml" ContentType='
                                    b'"application/vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml"/></Types>')
            elif item.filename == "xl/_rels/workbook.xml.rels":
                data = data.replace(b"</Relationships>", b'<Relationship Id="rIdMeta1" Type='
                                    b'"http://schemas.openxmlformats.org/officeDocument/2006/relationships/sheetMetadata" '
                                    b'Target="metadata.xml"/></Relationships>')
            zout.writestr(item, data)
        zout.writestr("xl/metadata.xml", DYN_METADATA)
    tmp.replace(path)


def load_yaml(path):
    import yaml
    with open(path, encoding="utf-8-sig") as f:
        return yaml.safe_load(f)


def main(argv=None, profile=None):
    ap = argparse.ArgumentParser(description="Build Jet Reports workbooks from a YAML spec.")
    ap.add_argument("--check", action="store_true", help="only check the spec")
    ap.add_argument("--profile", help="YAML profile (language, company key, colors)")
    ap.add_argument("spec")
    ap.add_argument("out", nargs="?")
    a = ap.parse_args(argv)
    if a.profile:
        profile = load_yaml(a.profile)
    spec = load_yaml(a.spec)
    if a.check:
        errs = check(spec, profile)
        print("OK - no findings" if not errs else "\n".join(errs))
        return 1 if errs else 0
    if not a.out:
        ap.error("out is required unless --check")
    try:
        print("OK:", build(spec, a.out, profile))
        return 0
    except ValueError as e:
        print(e)
        return 1


if __name__ == "__main__":
    sys.exit(main())
