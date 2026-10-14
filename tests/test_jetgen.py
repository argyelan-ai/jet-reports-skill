# -*- coding: utf-8 -*-
import re
import sys
import zipfile
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "skills" / "jet-reports" / "scripts"))
import jetgen  # noqa: E402

GERMAN = {
    "keywords": {"auto": "Auto+Ausblenden", "values": "+Werte", "option": "Option"},
    "options_sheet": "Optionen", "option_tags": ["Titel", "Wert", "Lookup", "Tooltip"],
    "company_key": "Mandant=", "require_table_numbers": True,
    "localized_function_names": ["WENN", "SUMME"],
}


def spec():
    return {
        "title": "Item list",
        "options": [
            {"name": "Company", "title": "Company", "value": "CRONUS",
             "lookup": 'NL("Lookup","2000000006 Company","1 Name")'},
            {"name": "MaxRows", "title": "Max rows", "value": 10},
        ],
        "sheets": [{
            "name": "Report",
            "title": '="Items - "&Company',
            "cells": [{"cell": "D4", "formula": 'NL("Count","27 Item","1 No.","Company=",Company)', "format": "int"}],
            "table": {
                "header_row": 6, "start_column": "C",
                "hide_row_if": "[[Inventory]]<0",
                "columns": [
                    {"header": "No.", "formula": 'NL("Rows","27 Item","1 No.","Company=",Company,"Limit=",MaxRows)'},
                    {"header": "Key", "hide": True,
                     "formula": 'NL("First","27 Item",,"Company=",Company,"1 No.","@@"&[[No.]])'},
                    {"header": "Description", "formula": 'NF([[Key]],"3 Description")'},
                    {"header": "Inventory", "formula": 'NF([[Key]],"68 Inventory")', "format": "int"},
                    {"header": "Status", "formula": 'IF([[Inventory]]>0,"stock","none")'},
                ],
            },
        }],
    }


def sheet_xml(path, title):
    z = zipfile.ZipFile(path)
    wb = z.read("xl/workbook.xml").decode()
    rels = z.read("xl/_rels/workbook.xml.rels").decode()
    rid = re.search(r'<sheet [^>]*name="%s"[^>]*r:id="(rId\d+)"' % re.escape(title), wb).group(1)
    rel = re.search(r'<Relationship [^>]*Id="%s"[^>]*>' % rid, rels).group(0)
    return z.read("xl/" + re.search(r'Target="/?(?:xl/)?([^"]+)"', rel).group(1)).decode()


def test_dynamic_array_jet_formulas(tmp_path):
    out = jetgen.build(spec(), tmp_path / "t.xlsx")
    x = sheet_xml(out, "Report")
    assert re.search(r'<c r="C7"[^>]*cm="1"><f t="array" ref="C7">_xll\.NL\("Rows"', x)
    assert '_xll.NL("First","27 Item",,"Company=",Company,"1 No.","@@"&amp;$C7)' in x
    assert '_xll.NF($D7,"3 Description")' in x
    assert re.search(r'<c r="G7"[^>]*><f>IF\(\$F7&gt;0', x)          # plain Excel formula stays plain
    assert "xl/metadata.xml" in zipfile.ZipFile(out).namelist()


def test_keywords_default_english(tmp_path):
    wb = openpyxl.load_workbook(jetgen.build(spec(), tmp_path / "t.xlsx"))
    assert wb.sheetnames == ["Options", "Report", "Version"]
    assert wb["Options"]["A1"].value == "Auto+Hide+HideSheet"
    assert wb["Options"]["A3"].value == "Option"
    assert wb.defined_names["Company"].attr_text == "Options!$D$3"
    r = wb["Report"]
    assert r["A1"].value == "Auto+Hide" and r["D1"].value == "Hide"
    assert r["A7"].value == '=IF($F7<0,"Hide","Show")'


def test_values_mode(tmp_path):
    s = spec()
    s["mode"] = "values"
    assert openpyxl.load_workbook(jetgen.build(s, tmp_path / "t.xlsx"))["Report"]["A1"].value == "Auto+Hide+Values"


def test_german_profile(tmp_path):
    s = spec()
    for col in s["sheets"][0]["table"]["columns"]:
        col["formula"] = col["formula"].replace("Company=", "Mandant=")
    s["sheets"][0]["cells"][0]["formula"] = s["sheets"][0]["cells"][0]["formula"].replace("Company=", "Mandant=")
    wb = openpyxl.load_workbook(jetgen.build(s, tmp_path / "t.xlsx", GERMAN))
    assert wb.sheetnames[0] == "Optionen" and wb["Optionen"]["A1"].value == "Auto+Ausblenden+HideSheet"
    assert [wb["Optionen"][c + "1"].value for c in "CDEF"] == ["Titel", "Wert", "Lookup", "Tooltip"]
    assert wb["Report"]["A1"].value == "Auto+Ausblenden"


@pytest.mark.parametrize("formula,msg,profile", [
    ('NL("Rows","27 Item","1 No.")', "every query needs a company", None),
    ('NL("Rows","27 Item","1 No.","Company=",Company', "unbalanced", None),
    ('NL("Rows","27 Item","1 No.","Company=",Compny)', "Compny", None),
    ('NF([[Missing]],"3 Description")', "Missing", None),
    ('NL("Rows","Item","1 No.","Company=",Company)', "table number", {"require_table_numbers": True}),
    ('WENN([[No.]]="",1,2)', "English function names", GERMAN),
])
def test_check_finds_errors(formula, msg, profile):
    s = spec()
    s["sheets"][0]["table"]["columns"].append({"header": "X", "formula": formula})
    if profile is GERMAN:
        profile = dict(GERMAN, company_keys_accepted=["Company="])
    errs = jetgen.check(s, profile)
    assert any(msg in e for e in errs), errs


def test_check_clean():
    assert jetgen.check(spec()) == []


def test_build_refuses_bad_spec(tmp_path):
    s = spec()
    s["sheets"][0]["table"]["columns"][0]["formula"] = 'NL("Rows","27 Item","1 No.")'
    with pytest.raises(ValueError):
        jetgen.build(s, tmp_path / "t.xlsx")
